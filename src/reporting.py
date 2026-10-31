from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill


def _read_for_date(path, business_date):
    frame = pd.read_csv(path)
    return frame[frame["business_date"].astype(str) == business_date]


def _format_workbook(workbook):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    color_scale = ColorScaleRule(
        start_type="min",
        start_color="F8696B",
        end_type="max",
        end_color="63BE7B",
    )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 35)
        if sheet.max_row > 2:
            sheet.conditional_formatting.add(f"B2:B{sheet.max_row}", color_scale)


def create_report(cfg, business_date=None):
    root = Path(cfg["root"])
    output = root / cfg["paths"]["processed"]
    report_dir = root / cfg["paths"]["reports"]
    report_dir.mkdir(parents=True, exist_ok=True)

    all_summary = pd.read_csv(output / "portfolio_risk.csv")
    date = business_date or str(all_summary["business_date"].max())
    if date not in set(all_summary["business_date"].astype(str)):
        raise ValueError(f"No analytics are available for {date}")

    summary = _read_for_date(output / "portfolio_risk.csv", date)
    risk = _read_for_date(output / "risk_metric.csv", date)
    pnl = _read_for_date(output / "pnl_nav.csv", date)
    stress = _read_for_date(output / "stress_results.csv", date)
    exceptions = pd.read_csv(output / "dq_exceptions.csv")
    krd = risk.groupby("portfolio_id")[
        ["krd_2y", "krd_5y", "krd_10y", "krd_30y", "cs01"]
    ].sum().reset_index()

    report = report_dir / f"fixed_income_risk_{date}.xlsx"
    with pd.ExcelWriter(report, engine="openpyxl") as writer:
        pd.DataFrame(
            {"Notice": [cfg["project"]["synthetic_label"]], "As of": [date]}
        ).to_excel(writer, sheet_name="Portfolio Summary", index=False)
        summary.to_excel(
            writer, sheet_name="Portfolio Summary", index=False, startrow=3
        )
        krd.to_excel(writer, sheet_name="DV01 CS01", index=False)
        pnl.to_excel(writer, sheet_name="PnL Attribution", index=False)
        summary[
            [
                "portfolio_id",
                "var_95",
                "var_99",
                "pnl_95_exception",
                "pnl_99_exception",
            ]
        ].to_excel(writer, sheet_name="VaR Monitoring", index=False)
        stress.to_excel(writer, sheet_name="Stress Tests", index=False)
        exceptions.to_excel(writer, sheet_name="Data Quality", index=False)
        _format_workbook(writer.book)
    return report

